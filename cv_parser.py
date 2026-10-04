"""CV text extraction (PDF, DOCX, TXT, RTF). Files are processed in memory only."""
from __future__ import annotations

import io
import logging
import re
from pathlib import Path

from sh_app.utils.text import clean_whitespace

logger = logging.getLogger(__name__)

SUPPORTED = {".pdf", ".docx", ".txt", ".rtf"}
MAX_BYTES = 8 * 1024 * 1024
MIN_CHARS = 150


class CVParseError(ValueError):
    """User-facing CV reading error."""


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise CVParseError("This PDF is password-protected. Please upload an unlocked copy.")
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _rtf(data: bytes) -> str:
    s = data.decode("latin-1", errors="ignore")
    s = re.sub(r"\{\\\*[^{}]*\}", " ", s)
    s = re.sub(r"\\par[d]?", "\n", s)
    s = re.sub(r"\\'[0-9a-fA-F]{2}", " ", s)
    s = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", s)
    return re.sub(r"[{}]", "", s)


def parse_cv(data: bytes, filename: str) -> str:
    """Return cleaned CV text or raise :class:`CVParseError` with a helpful message."""
    ext = Path(filename or "").suffix.lower()
    if ext == ".doc":
        raise CVParseError("Legacy .doc files are not supported. Please save the CV as DOCX or PDF.")
    if ext not in SUPPORTED:
        raise CVParseError("Unsupported file type. Please upload a PDF, DOCX, TXT or RTF file.")
    if not data:
        raise CVParseError("The uploaded file is empty.")
    if len(data) > MAX_BYTES:
        raise CVParseError("The file is larger than 8 MB. Please upload a smaller CV.")
    try:
        if ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        elif ext == ".rtf":
            text = _rtf(data)
        else:
            text = data.decode("utf-8", errors="replace")
    except CVParseError:
        raise
    except Exception as exc:
        logger.warning("CV parsing failed: %s", type(exc).__name__)
        raise CVParseError("Could not read this file. It may be corrupted or not a valid CV.") from exc
    text = clean_whitespace(text)
    if len(text) < MIN_CHARS:
        raise CVParseError(
            "Very little text could be extracted. If this is a scanned image PDF, "
            "please upload a text-based PDF or DOCX."
        )
    return text
